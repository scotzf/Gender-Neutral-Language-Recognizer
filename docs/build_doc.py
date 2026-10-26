import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
REPO = sys.argv[1] if len(sys.argv) > 1 else "[Insert GitHub link here]"
d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
def center(t, size=12, bold=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)
center("PROJECT IN AUTOMATA AND LANGUAGE THEORY", 14, True)
center("GENDER NEUTRAL LANGUAGE RECOGNIZER", 18, True)
center("\nSubmitted by:"); center("Denver Ugay", 12, True); center("BSCS 3B")
center("\nSubmitted to:"); center("ALMA MAE B. AUXTERO", 12, True); center("Instructor")
center("\nDate submitted: October 26, 2026")
d.add_page_break()

d.add_heading("1 Project overview", 1)
d.add_paragraph("NeutralScan FSA is a web application that reads a text document, finds predefined gender-specific words and phrases, and suggests gender-neutral alternatives. It shows the detected terms in a results table with the number of occurrences and suggestions, highlights the matches in the text, and displays the automaton states used to recognize a selected term. Suggestions are for review only; no automatic replacement is done.")

d.add_heading("2 Tools and recognition approach", 1)
d.add_paragraph("Language and tools: HTML, CSS and JavaScript (runs in any modern browser, no installation). Optional readers: mammoth.js (.docx) and pdf.js (.pdf).")
p = d.add_paragraph(); p.add_run("Approach: Finite State Automaton (trie-based DFA).").bold = True
d.add_paragraph("All terms in the dictionary are inserted into a trie. Each node is a state, each character is a transition, and a node where a term ends is an accepting state. State q0 is the start state. Formally, the DFA is M = (Q, Σ, δ, q0, F) where Q is the set of trie nodes, Σ is the set of lowercase characters, δ(q, c) moves to the child of q labelled c, q0 is the root, and F is the set of nodes that end a term.")
d.add_paragraph("Ignoring capitalization: every input character is converted to lowercase before the transition (text[i].toLowerCase()).")
d.add_paragraph("Whole-word matching: a scan may only start where the previous character is not a letter or digit, and an accepting state is reported only if the next character is not a letter or digit. This is why \"he\" is not detected in \"the\" (the scan cannot start at the \"h\" of \"the\" because the previous character \"t\" is a letter).")
d.add_paragraph("Sample transition table for the term \"fireman\" (all other inputs go to the dead state and the scan restarts):")
states = ["q0"]+[f"q{i}" for i in range(1,8)]
t = d.add_table(rows=2, cols=8); t.style = 'Table Grid'
word = "fireman"
t.rows[0].cells[0].text = "State"; t.rows[1].cells[0].text = "Next on →"
for i in range(7):
    t.rows[0].cells[i+1].text = f"q{i}"
    t.rows[1].cells[i+1].text = f"'{word[i]}' → q{i+1}" + (" (accept)" if i == 6 else "")
d.add_paragraph("Reading \"Fireman\": q0 -f→ q1 -i→ q2 -r→ q3 -e→ q4 -m→ q5 -a→ q6 -n→ q7. q7 is accepting and the next character is a space, so the term is recognized.")
d.add_paragraph("Predefined terms and suggested alternatives:")
terms = [("he","they"),("she","they"),("him","them"),("her","them/their"),("chairman","chairperson"),("stewardess","flight attendant"),("policeman","police officer"),("fireman","firefighter"),("manpower","workforce"),("mankind","humankind"),("ladies and gentlemen","everyone"),("salesman","salesperson"),("businessman","businessperson"),("spokesman","spokesperson"),("mailman","mail carrier"),("waitress","server")]
t = d.add_table(rows=1, cols=2); t.style = 'Table Grid'
t.rows[0].cells[0].text = "Term"; t.rows[0].cells[1].text = "Suggested alternative"
for a, b in terms:
    r = t.add_row().cells; r[0].text = a; r[1].text = b

d.add_heading("3 Application screenshots", 1)
def shot(f, cap, w=6.3):
    d.add_picture("screenshots/"+f, width=Inches(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True
shot("start.png", "Figure 1. Application start screen with the file upload button.")
shot("sample.png", "Figure 2. sample.txt loaded: detected terms, counts and suggestions, with highlighted text.")
shot("none.png", "Figure 3. A document with no detected terms.")
d.add_paragraph("Source code / GitHub link: " + REPO)

d.add_heading("4 Testing", 1)
tests = [
 ("1. Document with listed terms (sample.txt)", "Terms detected with counts and suggestions", "chairman 3, policeman 1, fireman 1, stewardess 1, mankind 1, salesman 1, manpower 1, ladies and gentlemen 1, waitress 1, she 1 (12 matches, 10 unique)", "Pass"),
 ("2. Same term repeated with different capitalization", "\"The CHAIRMAN and the chairman met. Chairman Lee said the policeman was there.\" → chairman = 3, policeman = 1", "chairman 3 (chairperson), policeman 1 (police officer)", "Pass"),
 ("3. No listed terms, includes \"the\", \"then\", \"they\"", "\"No predefined gender-specific terms were detected.\"", "Message shown; \"he\" not detected inside \"the\"/\"then\"/\"they\"", "Pass"),
 ("4. Empty .txt file", "Clear message for an empty file", "\"The file is empty.\"", "Pass"),
 ("5. Unsupported format (.csv)", "Clear unsupported-format message", "\"Unsupported file format. Please upload a .txt, .docx or .pdf file.\"", "Pass"),
]
t = d.add_table(rows=1, cols=4); t.style = 'Table Grid'
for i, h in enumerate(["Test situation", "Expected result", "Actual result", "Passed?"]): t.rows[0].cells[i].text = h
for row in tests:
    c = t.add_row().cells
    for i, v in enumerate(row): c[i].text = v
d.add_paragraph()
shot("multi.png", "Figure 4. Test 2: repeated term with different capitalization.", 5.5)
shot("empty.png", "Figure 5. Test 4: empty file.", 5.5)
shot("unsup.png", "Figure 6. Test 5: unsupported file format.", 5.5)

d.add_heading("5 Conclusion", 1)
d.add_paragraph("I completed a working gender-neutral language recognizer that uses a trie-based finite automaton to detect predefined terms case-insensitively and as whole words, then reports counts and suggested alternatives. I learned how a dictionary of words can be converted into a DFA and how word-boundary checks prevent false matches such as \"he\" in \"the\". Limitations: the dictionary is fixed and small, pronouns such as \"he\" and \"she\" are flagged even when they refer to a specific person, plural or irregular forms (e.g. \"policemen\") are not included, and suggestions need human review because the right replacement depends on context.")

d.add_heading("6 Sources and contributions", 1)
d.add_paragraph("References: course notes on finite automata; mammoth.js (docx reading) and pdf.js (PDF reading) via cdnjs. AI assistance: Claude (Anthropic) was used to help build and test the application and draft this document.")
d.add_paragraph("Contribution: Denver Ugay (BSCS 3B) – individual project: design, implementation, testing and documentation.")
d.save("Gender_Neutral_Language_Recognizer_Documentation.docx")
